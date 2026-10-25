"""以既有任務單 docx 為範本，填入 4 題並輸出練習題、診斷題兩份任務單（yincai-physics-node-items 使用）。

用法：
  python scripts/build_node_item_sheet.py 範本.docx items.json outputs/節點_標題_P4_MMDD.docx

實際寫出 節點_標題_P4_MMDD_練習題.docx（理解題、應用題）與 節點_標題_P4_MMDD_診斷題.docx（診斷題1、2），
拆分規則見 split_node_item_sheet.py。範本須是含 4 題的合併任務單，不能用已拆分的單份。
需要 python-docx 與 lxml。範本的追蹤修訂會全部接受、註解會清除；
第一個含「題目：」的表格視為範例，保持不動。
"""
import copy, io, json, re, sys
from pathlib import Path
from docx import Document
from split_node_item_sheet import output_paths, split_docx_bytes
from docx.oxml.ns import qn
from lxml import etree

tpl, spec, out = sys.argv[1:4]
data = json.load(open(spec, encoding="utf-8"))
doc = Document(tpl)

# 移除註解與追蹤修訂標記（範本可能是之前的修訂檔）
body = doc.element.body
for tag in ("w:commentRangeStart", "w:commentRangeEnd"):
    for el in list(body.iter(qn(tag))):
        el.getparent().remove(el)
for ref in list(body.iter(qn("w:commentReference"))):
    ref.getparent().getparent().remove(ref.getparent())
for rid, rel in list(doc.part.rels.items()):
    if "comments" in rel.reltype.lower():
        doc.part.drop_rel(rid)
for d in list(body.iter(qn("w:del"))):
    d.getparent().remove(d)
for ins in list(body.iter(qn("w:ins"))):
    parent = ins.getparent()
    if parent is None:
        continue
    if parent.tag == qn("w:rPr"):
        parent.remove(ins); continue
    i = parent.index(ins)
    for c in list(ins):
        parent.insert(i, c); i += 1
    parent.remove(ins)


def set_cell(cell, lines):
    if isinstance(lines, str):
        lines = [lines]
    tc = cell._tc
    paras = tc.findall(qn("w:p"))
    has_text = lambda el: any((t.text or "").strip() for t in el.iter(qn("w:t")))
    tp = next((p for p in tc.iter(qn("w:p")) if has_text(p)), paras[0])
    tr = next((r for r in tp.iter(qn("w:r")) if has_text(r)), None)
    ppr = tp.find(qn("w:pPr"))
    rpr = tr.find(qn("w:rPr")) if tr is not None else None
    for child in list(tc):
        if child.tag in (qn("w:p"), qn("w:tbl")):
            tc.remove(child)
    for line in lines:
        p = etree.SubElement(tc, qn("w:p"))
        if ppr is not None:
            p.append(copy.deepcopy(ppr))
        r = etree.SubElement(p, qn("w:r"))
        if rpr is not None:
            r.append(copy.deepcopy(rpr))
        t = etree.SubElement(r, qn("w:t"))
        t.text = line
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")


# 表頭（學習內容編號、知識節點編號、影片標題…）
hdr = doc.tables[0]
for r in hdr.rows:
    lab = r.cells[0].text.strip().rstrip("：:")
    if lab in data.get("header", {}):
        set_cell(r.cells[-1], data["header"][lab])

# 題目表格：第 1 個含「題目：」的表格是範例，不動；之後依序是 4 題
KEY = {"標題：": "標題", "來源：": "來源", "題目：": "題目", "（A）": "A", "（B）": "B",
       "（C）": "C", "（D）": "D", "詳解：": "詳解", "答案：": "答案"}
qtables = [t for t in doc.tables if any(r.cells[0].text.strip() == "題目：" for r in t.rows)]
assert len(qtables) == 1 + len(data["items"]), (
    f"範本有 {len(qtables)-1} 個題目表格，items 有 {len(data['items'])} 題；範本要用含 4 題的合併任務單，不是 _練習題／_診斷題 拆分檔")
for t, item in zip(qtables[1:], data["items"]):
    for r in t.rows:
        k = KEY.get(r.cells[0].text.strip())
        if k and k in item:
            set_cell(r.cells[-1], item[k])

# 題號標題統一成「N. 理解題」「N. 診斷題1（經典題修改）」
for p in doc.paragraphs:
    joined = "".join(r.text for r in p.runs)
    m = re.match(r"^\s*(\d)\s*[\.．]\s*", joined)
    if not m or not re.search(r"理解|應用|診斷", joined):
        continue
    new = f"{m.group(1)}. " + joined[m.end():].replace("(", "（").replace(")", "）")
    p.runs[0].text = new
    for r in p.runs[1:]:
        r.text = ""

# 拆成練習題、診斷題兩份；題號在各自檔案內從 1 重編
buffer = io.BytesIO()
doc.save(buffer)
parts = split_docx_bytes(buffer.getvalue())
for kind, target in output_paths(Path(out)).items():
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_bytes(parts[kind])
    print("saved", target)
