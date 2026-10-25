"""把 WSQ 學生版 Markdown 轉成一頁 DOCX（physics-worksheet-generator 使用）。

用法：
  python scripts/build_wsq_docx.py outputs/節點_WSQ.md outputs/節點_單元_WSQ.docx

只轉第一個標題含「WSQ」的一級標題到下一個一級標題之間（學生版），格式見 audit_wsq.py。
需要 python-docx。中文微軟正黑體、英文與數字 Arial。
手寫空間：句中的＿＿轉成加長的底線空格；整行底線（S、Q 的作答行）轉成兩條作答線；□ 換成較大的 ☐。
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from audit_wsq import BLANK_LINE, sections, student_block

EAST_ASIA = "微軟正黑體"
LATIN = "Arial"
BLANK = re.compile(r"＿+")
ANSWER_LINES = 2


def parse(text: str) -> dict:
    """回傳 {"title", "header", "sections": [(標題, [行...])]}；不依賴 python-docx，可單獨測試。"""
    block = student_block(text)
    if not block:
        raise ValueError("找不到標題含「WSQ」的學生版")
    header = next((line.strip() for line in block[1:] if line.strip() and not line.startswith("## ")), "")
    parts = sections(block)
    return {"title": block[0][2:].strip(), "header": header, "sections": list(parts.items())}


def build(spec: dict, out: Path) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Cm(2))

    def paragraph(text: str, size: float, bold: bool = False, before: float = 0, after: float = 4):
        p = doc.add_paragraph()
        p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(before), Pt(after)
        p.paragraph_format.line_spacing = 1.25
        pos = 0
        for m in list(BLANK.finditer(text)) + [None]:
            end = m.start() if m else len(text)
            pieces = [(text[pos:end].replace("□", "☐"), False)]
            if m:
                # 用不換行空格：Word 不畫段尾一般空格的底線，「姓名：」後的空格會不見
                pieces.append((" " * max(16, len(m.group()) * 6), True))
                pos = m.end()
            for value, underline in pieces:
                if not value:
                    continue
                run = p.add_run(value)
                run.font.size, run.font.bold, run.font.name = Pt(size), bold, LATIN
                run.font.underline = underline
                run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), EAST_ASIA)
        return p

    def answer_line():
        p = doc.add_paragraph()
        p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(0), Pt(6)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        p.paragraph_format.line_spacing = Pt(28)
        # 連續兩段框線相同時 Word 會合併成一組，只畫 bottom；加 between 才會每行都有線
        border = OxmlElement("w:pBdr")
        for edge in ("bottom", "between"):
            line = OxmlElement(f"w:{edge}")
            for key, value in (("val", "single"), ("sz", "6"), ("space", "1"), ("color", "808080")):
                line.set(qn(f"w:{key}"), value)
            border.append(line)
        p._p.get_or_add_pPr().append(border)

    title = paragraph(spec["title"], 18, bold=True, after=6)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if spec["header"]:
        paragraph(spec["header"], 13, after=8)
    for heading, lines in spec["sections"]:
        paragraph(heading, 14, bold=True, before=10, after=4)
        for line in lines:
            if BLANK_LINE.match(line):
                for _ in range(ANSWER_LINES):
                    answer_line()
            else:
                paragraph(line, 13, after=6)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)


def main(argv: list[str] | None = None) -> int:
    args = argv if argv is not None else sys.argv[1:]
    if len(args) != 2:
        print(__doc__)
        return 2
    source, out = Path(args[0]), Path(args[1])
    build(parse(source.read_text(encoding="utf-8")), out)
    print("saved", out)
    return 0


if __name__ == "__main__":
    sys.exit(main())
